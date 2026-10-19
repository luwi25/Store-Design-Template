"""One-off: apply a manually-downloaded POS Sales Analysis workbook to the
published catalog for BOTH stores.

Unlike weekly_sync.py (which pulls from the UBEST API and runs on a
schedule for Home Express Basak only), this reads category/product
rankings directly out of the user-supplied Excel report -- ranked by
gross sales (PHP), matching that report's own methodology -- and applies
them to:

  Home Express Basak (homeBack): "Best Seller This Season" + "Top
  Product" groups only (10 slots). "Back to School Essentials" (the
  Stationery group) is left untouched.

  PesoPlus Mambaling (peso = Front, pesoBack = Back), touched here for
  the first time:
    - peso "Small Appliances" (6 slots): top 6 products by gross sales
      within the Small Appliances department.
    - peso "Best Sellers This Season" (4) + "Top Products" (5): one top
      product per category, from the report's top-10-categories list
      with Small-Appliances-department categories excluded (handled
      above) -- 9 categories remain, matching 4+5 exactly.
    - pesoBack's 7 department groups (27 slots total): same department
      mapping as Home Express Front, ranked by gross sales, using
      pesoBack's own per-group slot counts.

Run manually: python scripts/apply_sales_analysis_excel.py
"""
import glob
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import weekly_sync as ws  # reuses REPO, load_api_key, api_get_all_pages, placeholder_image_marker, title_case, run, run_tests

import openpyxl

PESO_BACK_GROUP_DEPTS = [
    ("Kitchen & Dining", 6, ["Kitchenware & Tableware", "Cookware"]),
    ("Home Organizers", 3, ["Home Storage & Organization", "Kitchen Storage & Organization"]),
    ("Home Textiles", 3, ["Home Textiles"]),
    ("Household Cleaning", 4, ["Cleaning & Laundry Supplies"]),
    ("Personal Care", 4, ["Hygiene", "Beauty & Fragrance"]),
    ("Hardware & Tools", 4, ["Hardware & Electrical", "Tools"]),
    ("Toys & Party", 3, ["玩具 Toys", "Seasonal & Party"]),
]
SMALL_APPLIANCES_DEPT = "Small Appliances"
STATIONERY_DEPT = "Stationery"


def split_en(s):
    return ws.split_bilingual(s)[0] if s else None


def load_raw_rows(wb, sheet_name):
    sheet = wb[sheet_name]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def aggregate_by_stockno(rows):
    agg = {}
    for r in rows:
        stockno = r.get("Stock No (货号)")
        if not stockno:
            continue
        a = agg.setdefault(stockno, {
            "stockno": stockno, "name": r.get("Product (商品)"),
            "department": split_en(r.get("Department (部门)")),
            "category": split_en(r.get("Category (类别)")),
            "amount": 0.0, "qty": 0.0,
        })
        a["amount"] += r.get("Sales Amount (销售额)") or 0
        a["qty"] += r.get("Sales Qty (销量)") or 0
    return agg


def top10_categories(cat_sheet, name_col):
    return [cat_sheet.cell(r, name_col).value for r in range(7, 17)]


def top_product_in_category(agg, category_en):
    pool = [a for a in agg.values() if a["category"] == category_en]
    return max(pool, key=lambda a: a["amount"]) if pool else None


def top_products_in_depts(agg, depts, count):
    pool = sorted((a for a in agg.values() if a["department"] in depts), key=lambda a: -a["amount"])
    return pool[:count]


def category_department(agg, category_en):
    for a in agg.values():
        if a["category"] == category_en:
            return a["department"]
    return None


def fetch_catalog(base_url, api_key, store_id):
    items, _ = ws.api_get_all_pages(base_url, api_key, f"/stores/{store_id}/products")
    return {p["stockno"]: p for p in items}


def upsert_with_prefix(repo, state, products_by_id, prefix, stockno, name, price, template_id):
    """Same contract as weekly_sync.upsert_product (reuse cached photo for a
    returning stockno, flat placeholder for a brand-new one -- the UBEST
    image CDN is still down), just with a caller-chosen id prefix so
    PesoPlus products ('sales-ppm-') never collide with Home Express's
    ('sales-heb-')."""
    pid = f"{prefix}-{re.sub(r'[^A-Za-z0-9]+', '-', stockno).strip('-').lower()}"
    existing = products_by_id.get(pid)
    if existing is not None:
        image = existing["image"]
        photo_status = "cached"
    else:
        image = ws.placeholder_image_marker(repo)
        photo_status = "placeholder"

    product = {
        "id": pid, "templateId": template_id, "name": ws.title_case(name),
        "price": price if price is not None else 0, "code": stockno, "image": image,
    }
    if existing is not None:
        existing.update(product)
    else:
        products_by_id[pid] = product
        state["products"].append(product)
    return pid, photo_status


def resolve_price(store_id, item):
    catalog_entry = CATALOG_BY_STORE.get(store_id, {}).get(item["stockno"])
    if catalog_entry and catalog_entry.get("retail_price") is not None:
        return catalog_entry["retail_price"]
    return round(item["amount"] / item["qty"], 2) if item["qty"] else 0


def apply_group(repo, state, products_by_id, prefix, store_id, slots, group_items, template_id, log_label):
    by_group_idx = defaultdict(list)
    for i, s in enumerate(slots):
        by_group_idx[s["group"]].append(i)
    for g, items in group_items.items():
        idxs = by_group_idx.get(g, [])
        for idx, item in zip(idxs, items):
            price = resolve_price(store_id, item)
            pid, photo_status = upsert_with_prefix(repo, state, products_by_id, prefix, item["stockno"], item["name"], price, template_id)
            slots[idx] = {"productId": pid, "group": g, "discountPercent": 0}
            tag = " [placeholder photo]" if photo_status == "placeholder" else ""
            ws.log(f"  {log_label} [group {g}] {item['name']} (P{item['amount']:.2f} gross){tag}")


CATALOG_BY_STORE = {}


def main():
    excel_path = glob.glob(r"C:\Users\Admin\Downloads\Sales Analysis*.xlsx")[0]
    ws.log(f"reading {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    cat_sheet = wb["Category & Brand 类别品牌"]

    he_rows = load_raw_rows(wb, "Home Express Basak 数据")
    pp_rows = load_raw_rows(wb, "Pesoplus Mambaling 数据")
    he_agg = aggregate_by_stockno(he_rows)
    pp_agg = aggregate_by_stockno(pp_rows)

    base_url, api_key = ws.load_api_key()
    global CATALOG_BY_STORE
    ws.log("fetching catalogs for both stores (for photos/prices)...")
    CATALOG_BY_STORE["home-express-basak"] = fetch_catalog(base_url, api_key, "home-express-basak")
    CATALOG_BY_STORE["pesoplus-mambaling"] = fetch_catalog(base_url, api_key, "pesoplus-mambaling")

    project_path = ws.REPO / "defaults" / "project.json"
    project = json.loads(project_path.read_text(encoding="utf-8"))
    state = project["state"]
    products_by_id = {p["id"]: p for p in state["products"]}

    # ---------- Home Express Basak: Best Seller This Season + Top Product ----------
    he_cats = [c for c in top10_categories(cat_sheet, 2) if c]
    he_picks = []
    for cat in he_cats:
        cat_en = split_en(cat)
        if category_department(he_agg, cat_en) == STATIONERY_DEPT:
            continue
        item = top_product_in_category(he_agg, cat_en)
        if item:
            he_picks.append(item)
    assert len(he_picks) >= 10, f"expected >=10 Home Express category picks, got {len(he_picks)}"
    homeback_slots = state["homeBack"]["slots"]
    apply_group(ws.REPO, state, products_by_id, "sales-heb", "home-express-basak", homeback_slots,
                {1: he_picks[:5], 2: he_picks[5:10]}, "home-back", "HE BACK")

    # ---------- PesoPlus Mambaling: Front (peso) ----------
    pp_cats = [c for c in top10_categories(cat_sheet, 8) if c]
    appliance_cats, other_cats = [], []
    for cat in pp_cats:
        cat_en = split_en(cat)
        (appliance_cats if category_department(pp_agg, cat_en) == SMALL_APPLIANCES_DEPT else other_cats).append(cat_en)
    ws.log(f"PesoPlus top categories: {len(other_cats)} general, {len(appliance_cats)} small-appliance")
    assert len(other_cats) >= 9, f"expected >=9 non-appliance PesoPlus categories, got {len(other_cats)}"

    general_picks = [top_product_in_category(pp_agg, c) for c in other_cats]
    general_picks = [p for p in general_picks if p][:9]
    appliance_picks = top_products_in_depts(pp_agg, [SMALL_APPLIANCES_DEPT], 6)

    peso_slots = state["peso"]["slots"]
    apply_group(ws.REPO, state, products_by_id, "sales-ppm", "pesoplus-mambaling", peso_slots,
                {0: general_picks[:4], 1: general_picks[4:9], 2: appliance_picks}, "peso-plus", "PESO FRONT")

    # ---------- PesoPlus Mambaling: Back (pesoBack) ----------
    pesoback_slots = state["pesoBack"]["slots"]
    dept_group_items = {}
    for g, (label, count, depts) in enumerate(PESO_BACK_GROUP_DEPTS):
        picks = top_products_in_depts(pp_agg, depts, count)
        ws.log(f"PESO BACK group {g} '{label}': picked={len(picks)}")
        dept_group_items[g] = picks
    apply_group(ws.REPO, state, products_by_id, "sales-ppm", "pesoplus-mambaling", pesoback_slots,
                dept_group_items, "peso-back", "PESO BACK")

    # ---------- prune orphaned sales-driven products from earlier runs ----------
    referenced = set()
    for key in ("everyday", "homeBack", "peso", "pesoBack"):
        referenced |= {s["productId"] for s in state[key]["slots"]}
    kept = [p for p in state["products"]
            if not (p["id"].startswith("sales-heb-") or p["id"].startswith("sales-ppm-")) or p["id"] in referenced]
    pruned = len(state["products"]) - len(kept)
    state["products"] = kept
    if pruned:
        ws.log(f"pruned {pruned} orphaned sales-driven product(s)")

    project["sourceRevision"] += 1
    project_path.write_text(json.dumps(project, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    ws.log(f"wrote defaults/project.json, sourceRevision -> {project['sourceRevision']}, products: {len(state['products'])}")

    if not ws.run_tests(ws.REPO):
        ws.flush_log(False, "Tests failed. Nothing committed/pushed; working tree has the attempted change.")
        ws.log("ABORTED before push: tests failed.")
        sys.exit(1)

    ws.run(["git", "add", "defaults"], ws.REPO)
    import subprocess
    diff_check = subprocess.run(["git", "diff", "--cached", "--quiet"], cwd=str(ws.REPO))
    if diff_check.returncode == 0:
        ws.log("nothing to commit.")
        ws.flush_log(True)
        return
    ws.run(["git", "commit", "-m", "Apply POS Sales Analysis report (2026-10-12 to 2026-10-18) to both stores"], ws.REPO)
    if ws.run(["git", "push", "origin", "main"], ws.REPO) != 0:
        ws.flush_log(False, "git push failed; commit is local only.")
        ws.log("ABORTED: push failed.")
        sys.exit(1)
    ws.log("done: pushed to GitHub.")
    ws.flush_log(True)


if __name__ == "__main__":
    main()
